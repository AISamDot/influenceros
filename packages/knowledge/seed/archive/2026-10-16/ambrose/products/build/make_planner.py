"""Builds the Planner Sheet (planner-sheet.xlsx) for The Butler's Holiday Plan.
Run: python3 make_planner.py [out.xlsx]
The same file is uploaded to Google Drive with conversion to make the Google template (build-spec.md, Job J2)."""
import sys, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import LineChart, Reference

OUT = sys.argv[1] if len(sys.argv) > 1 else "planner-sheet.xlsx"
INK, MID, SILVER, LINEN, AMBER = "1F3440", "161D2A", "C9CCCF", "F1EEE7", "D9A55B"
HDR = PatternFill("solid", fgColor=INK); INP = PatternFill("solid", fgColor=LINEN)
AMB = PatternFill("solid", fgColor=AMBER)
thin = Side(style="thin", color=SILVER); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
TICK = "✓"
SHARES = ["Showpiece", "Sides and bread", "Nibbles or starter (appetizers)", "Pudding (dessert)",
          "Drinks (incl. non-alcoholic, ice)", "The table"]
PCT = [0.35, 0.25, 0.10, 0.10, 0.15, 0.05]

wb = Workbook()
def sheet(name, title, note=None, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = name
    ws["A1"] = title; ws["A1"].font = Font(bold=True, size=14, color=INK)
    if note: ws["A2"] = note; ws["A2"].font = Font(italic=True, color="3A3B3D")
    ws.sheet_view.showGridLines = False
    return ws
def header(ws, row, labels, col=1):
    for i, l in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=l)
        c.fill = HDR; c.font = Font(bold=True, color="FFFFFF"); c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BOX
def inp(c, fmt=None):
    c.fill = INP; c.border = BOX
    if fmt: c.number_format = fmt
def out(c, fmt=None, bold=False):
    c.border = BOX; c.font = Font(bold=bold, color=MID)
    if fmt: c.number_format = fmt
def widths(ws, ws_w):
    for k, v in ws_w.items(): ws.column_dimensions[k].width = v
def ticklist(ws, rng):
    dv = DataValidation(type="list", formula1=f'"{TICK}"', allow_blank=True); ws.add_data_validation(dv); dv.add(rng)
def listval(ws, rng, items):
    dv = DataValidation(type="list", formula1='"' + ",".join(items) + '"', allow_blank=True); ws.add_data_validation(dv); dv.add(rng)

USD = '"$"#,##0.00'; DATE = "ddd d mmm yyyy"; TIME = "h:mm AM/PM"; DT = "ddd d mmm yyyy h:mm AM/PM"

# 1 At a Glance
g = sheet("At a Glance", "Your Gathering at a Glance", "Type in the Linen cells only. Every other tab reads from this one. Thanksgiving 2026: Thu 26 Nov 2026 · Christmas: Fri 25 Dec 2026.", first=True)
rows = [("Occasion", None, None), ("T-0 date (the day itself)", None, DATE), ("Doorbell time", None, TIME),
        ("Sit-down time", None, TIME), ("End time", None, TIME), ("Guests, host included", None, "0"),
        ("Total budget: food, drink, table", None, USD)]
for i, (lab, v, f) in enumerate(rows, start=4):
    g.cell(row=i, column=1, value=lab); inp(g.cell(row=i, column=2, value=v), f)
calc = [("Per-Plate Number (budget ÷ guests)", '=IFERROR(B10/B9,"")', USD),
        ("Hours, doorbell to end", '=IF(OR(B6="",B8=""),"",ROUND(MOD(B8-B6,1)*24,1))', "0.0"),
        ("T-0 date the Clock is using", "=IF(Reckoning!B27<>\"\",Reckoning!B27,B5)", DATE),
        ("Today", "=TODAY()", DATE),
        ("Today's T-number", '=IF(B13="","",IF(B13-B14>0,"T-"&(B13-B14),IF(B13=B14,"T-0","T+"&(B14-B13))))', None)]
for i, (lab, f, nf) in enumerate(calc, start=11):
    g.cell(row=i, column=1, value=lab).font = Font(bold=True); out(g.cell(row=i, column=2, value=f), nf, bold=True)
widths(g, {"A": 40, "B": 26})

# 13 Reckoning is referenced above at B27 (next year's T-0 date); built last.

# 2 Clock
c = sheet("Clock", "The Clock: 25 tasks, dated from your T-0", "Dates come from 'At a Glance' (or next year's date on 'Reckoning', once you type it). Tick with ✓.")
c["A3"] = "Today's task:"; c["A3"].font = Font(bold=True)
c["B3"] = '=IF(\'At a Glance\'!B13="","Type your T-0 date on At a Glance",IFERROR(INDEX(D6:D30,MATCH(TODAY(),C6:C30,0)),"No task dated today: see the next date below"))'
c["F3"] = "Ticks:"; c["G3"] = '=COUNTIF(G6:G30,"' + TICK + '")&" / 25"'
c["F4"] = "Minutes:"; c["G4"] = '=COUNTIF(G6:G30,"' + TICK + '")*15'
header(c, 5, ["T", "Offset", "Date", "Today's fifteen minutes", "Page", "Write down one number", "Done", "My number"])
TASKS = [(-21, "Fill in At a Glance; write the guest list; Briefing 1", 5, "Guests, host included"),
 (-20, "Set the Per-Plate Number", 12, "Per plate $"), (-19, "The Count: chairs, plates, glasses, napkins, dishes", 21, "Chairs short"),
 (-18, "One Showpiece; make-ahead menu around it", 27, "Dishes on the menu"), (-17, "Oven Timetable, first draft", 30, "Oven dishes in the last hour"),
 (-16, "Gift Ceiling: season total, ceiling per name", 15, "Season gift total $"), (-15, "Staff & Service list; count the week's ticks", 16, "Staff & Service total $"),
 (-14, "Send the invitation; Briefing 2", 22, "Invitations sent"), (-13, "Assign dishes, not dollars", 14, "Dishes assigned"),
 (-12, "Choose the table level; draft the seating", "18, 26", "Place settings"), (-11, "Drinks maths", 28, "Bottles needed"),
 (-10, "Shopping list, part one (Christmas: Staff & Service deadline)", "28, 16", "Items on the list"), (-9, "RSVP chase; recalculate per-plate", "22, 12", "Confirmed guests"),
 (-8, "Dietary needs", 23, "Dietary notes"), (-7, "Write the Guest Brief; one-week note; thaw sum; Briefing 3", "22, 29", "Thaw starts on"),
 (-6, "Borrow chairs, not money", 21, "Items borrowed"), (-5, "Shopping, part one; log the spend", "13, 28", "Spent so far $"),
 (-4, "Clear the route: coats, loo (bathroom), hall", 32, "Coat hooks free"), (-3, "Make-ahead plan, final; check the thaw", "30, 29", "Dishes made ahead"),
 (-2, "Send the Guest Brief; fresh shopping", "22, 13", "Spent so far $"), (-1, "Set the Night Before; Briefing 4", 33, "Items ticked / 16"),
 (0, "Run-of-show; the 10-Minute Sweep; sit down", "34, 35", "Calm score at the door"), (1, "Clear-down; borrowed things returned", "36, 38", "Borrowed items returned"),
 (2, "Thank-you notes; Briefing 5", 37, "Notes sent"), (7, "The Reckoning", 39, "Actual per plate $")]
for i, (off, task, page, num) in enumerate(TASKS, start=6):
    c.cell(row=i, column=1, value=("T" + ("+" if off > 0 else ("-" if off < 0 else "-")) + str(abs(off)))).border = BOX
    c.cell(row=i, column=2, value=off).border = BOX
    out(c.cell(row=i, column=3, value=f"=IF('At a Glance'!$B$13=\"\",\"\",'At a Glance'!$B$13+B{i})"), DATE)
    c.cell(row=i, column=4, value=task).border = BOX
    c.cell(row=i, column=5, value=str(page)).border = BOX
    c.cell(row=i, column=6, value=num).border = BOX
    inp(c.cell(row=i, column=7)); inp(c.cell(row=i, column=8))
ticklist(c, "G6:G30")
c.conditional_formatting.add("A6:H30", FormulaRule(formula=["$C6=TODAY()"], fill=AMB))
widths(c, {"A": 7, "B": 7, "C": 18, "D": 52, "E": 8, "F": 26, "G": 7, "H": 16})

# 3 Per-Plate
p = sheet("Per-Plate", "Where the plate goes", "Change the percentages to suit your menu; they must add up to 100%.")
header(p, 4, ["Share", "%", "Amount", "Per plate"])
for i, (s, pc) in enumerate(zip(SHARES, PCT), start=5):
    p.cell(row=i, column=1, value=s).border = BOX
    inp(p.cell(row=i, column=2, value=pc), "0%")
    out(p.cell(row=i, column=3, value=f"=IF('At a Glance'!$B$10=\"\",\"\",'At a Glance'!$B$10*B{i})"), USD)
    out(p.cell(row=i, column=4, value=f"=IFERROR(C{i}/'At a Glance'!$B$9,\"\")"), USD)
p["A11"] = "Total"; p["A11"].font = Font(bold=True)
out(p["B11"], "0%", True); p["B11"] = "=SUM(B5:B10)"
out(p["C11"], USD, True); p["C11"] = '=IF(\'At a Glance\'!B10="","",SUM(C5:C10))'
p["A12"] = '=IF(ROUND(B11,4)=1,"Adds up to 100%. Good.","Adjust: the shares add up to "&TEXT(B11,"0%"))'
widths(p, {"A": 36, "B": 8, "C": 14, "D": 12})

# 4 Spend Log
s = sheet("Spend Log", "The Spend Log", "Every receipt, the same evening. A share turns amber when it passes its line.")
header(s, 4, ["Date", "Where", "What", "Share", "Planned $", "Actual $"])
for r in range(5, 55):
    for col in range(1, 7): inp(s.cell(row=r, column=col))
    s.cell(row=r, column=1).number_format = DATE; s.cell(row=r, column=5).number_format = USD; s.cell(row=r, column=6).number_format = USD
listval(s, "D5:D54", SHARES)
header(s, 4, ["Share", "Line (Per-Plate)", "Spent", "Left"], col=8)
for i, sh in enumerate(SHARES, start=5):
    s.cell(row=i, column=8, value=sh).border = BOX
    out(s.cell(row=i, column=9, value=f"='Per-Plate'!C{i}"), USD)
    out(s.cell(row=i, column=10, value=f'=SUMIF($D$5:$D$54,H{i},$F$5:$F$54)'), USD)
    out(s.cell(row=i, column=11, value=f'=IF(I{i}="","",I{i}-J{i})'), USD)
s.conditional_formatting.add("H5:K10", FormulaRule(formula=['AND(ISNUMBER($I5),$J5>$I5)'], fill=AMB))
s["H12"] = "Total spent"; out(s["J12"], USD, True); s["J12"] = "=SUM(F5:F54)"
s["H13"] = "Per plate so far"; out(s["J13"], USD, True); s["J13"] = "=IFERROR(J12/'At a Glance'!B9,\"\")"
s["H14"] = "Within the number?"; out(s["J14"], None, True); s["J14"] = "=IF('At a Glance'!B10=\"\",\"\",IF(J12<='At a Glance'!B10,\"Yes\",\"No: over by $\"&TEXT(J12-'At a Glance'!B10,\"0.00\")))"
widths(s, {"A": 16, "B": 16, "C": 24, "D": 28, "E": 11, "F": 11, "G": 3, "H": 30, "I": 15, "J": 12, "K": 12})

# 11 Oven (built before Assignments formulas reference it by name; order fixed later)
o = wb.create_sheet("Oven"); o.sheet_view.showGridLines = False

# 5 Assignments
a = sheet("Assignments", "Assign dishes, not dollars", "If a guest's dish needs your oven, give the times; clashes with the Oven tab are flagged.")
header(a, 4, ["Guest", "Their one job", "Serves", "Arrives in (their dish / yours)", "Needs", "Oven from", "Oven to", "Confirmed", "Clash?"])
for r in range(5, 25):
    for col in range(1, 9): inp(a.cell(row=r, column=col))
    a.cell(row=r, column=6).number_format = TIME; a.cell(row=r, column=7).number_format = TIME
    out(a.cell(row=r, column=9, value=f'=IF(AND(E{r}="Oven",F{r}<>"",G{r}<>""),IF(SUMPRODUCT((Oven!$B$5:$B$34="Oven")*(Oven!$D$5:$D$34<G{r})*(Oven!$E$5:$E$34>F{r}))>0,"Clash with the Oven tab",""),"")'))
listval(a, "E5:E24", ["Oven", "Fridge", "Neither"]); ticklist(a, "H5:H24")
a.conditional_formatting.add("I5:I24", FormulaRule(formula=['$I5<>""'], fill=AMB))
widths(a, {"A": 16, "B": 26, "C": 8, "D": 22, "E": 10, "F": 11, "G": 11, "H": 10, "I": 24})

# 6 Drinks
d = sheet("Drinks", "Drinks maths", "A planning sum for buying, not a target for drinking. 750 ml bottle ≈ five 150 ml (5 oz) glasses.")
dl = [("A · Guests who'll drink alcohol", None, True, "0"),
      ("B · Hours (from At a Glance; overwrite if you like)", "='At a Glance'!B12", True, "0.0"),
      ("C · Drinks per drinking guest: 2 + (B − 1)", '=IF(B5="","",2+(B5-1))', False, "0.0"),
      ("D · Total alcoholic drinks: A × C", '=IF(OR(B4="",B6=""),"",B4*B6)', False, "0"),
      ("E · Bottles of wine, if all wine: D ÷ 5, rounded up", '=IF(B7="","",ROUNDUP(B7/5,0))', False, "0"),
      ("F · Guests not drinking alcohol", None, True, "0"),
      ("Spare non-alcoholic drinks for drivers", 3, True, "0"),
      ("G · Non-alcoholic drinks: F × C + spare", '=IF(OR(B9="",B6=""),"",B9*B6+B10)', False, "0"),
      ("H · Ice, lb (about 1 lb per guest)", "=IF('At a Glance'!B9=\"\",\"\",'At a Glance'!B9*1)", False, "0"),
      ("H · Ice, kg (about 0.5 kg per guest)", "=IF('At a Glance'!B9=\"\",\"\",'At a Glance'!B9*0.5)", False, "0.0"),
      ("I · Water jugs (one per four guests)", "=IF('At a Glance'!B9=\"\",\"\",ROUNDUP('At a Glance'!B9/4,0))", False, "0")]
for i, (lab, f, is_in, nf) in enumerate(dl, start=4):
    d.cell(row=i, column=1, value=lab); cell = d.cell(row=i, column=2, value=f)
    (inp if is_in else out)(cell, nf)
widths(d, {"A": 50, "B": 14})

# 7 Shopping
sh = sheet("Shopping", "The Shopping List", "Dry goods at T-5; fresh things at T-2. Totals check against your Per-Plate shares.")
header(sh, 4, ["Item", "Qty", "Share", "Est. $", "When", "Bought", "Assigned to a guest instead?"])
for r in range(5, 65):
    for col in range(1, 8): inp(sh.cell(row=r, column=col))
    sh.cell(row=r, column=4).number_format = USD
listval(sh, "C5:C64", SHARES); listval(sh, "E5:E64", ["T-5", "T-2"]); ticklist(sh, "F5:F64")
header(sh, 4, ["Share", "Line", "Est. on list", "Left"], col=9)
for i, shn in enumerate(SHARES, start=5):
    sh.cell(row=i, column=9, value=shn).border = BOX
    out(sh.cell(row=i, column=10, value=f"='Per-Plate'!C{i}"), USD)
    out(sh.cell(row=i, column=11, value=f"=SUMIF($C$5:$C$64,I{i},$D$5:$D$64)"), USD)
    out(sh.cell(row=i, column=12, value=f'=IF(J{i}="","",J{i}-K{i})'), USD)
sh.conditional_formatting.add("I5:L10", FormulaRule(formula=['AND(ISNUMBER($J5),$K5>$J5)'], fill=AMB))
sh["I12"] = "T-5 items"; sh["K12"] = '=COUNTIF(E5:E64,"T-5")'
sh["I13"] = "T-2 items"; sh["K13"] = '=COUNTIF(E5:E64,"T-2")'
sh["I14"] = "Bought"; sh["K14"] = '=COUNTIF(F5:F64,"' + TICK + '")'
widths(sh, {"A": 26, "B": 6, "C": 28, "D": 10, "E": 7, "F": 8, "G": 24, "H": 3, "I": 30, "J": 12, "K": 12, "L": 12})

# 8 Gift Ledger (Staff & Service total arrives from its own tab)
gl = sheet("Gift Ledger", "The Gift Ledger", "A ceiling per name, in ink. Use it, Eat it, or Keep it ten years. A row turns amber when Spent passes its ceiling.")
labs = [("Season gift total (your limit, incl. Staff & Service)", None, True), ("Staff & Service total (from its tab)", "='Staff & Service'!C17", False),
        ("Gift ceilings, all names", "=SUM(B12:B51)", False), ("Ceilings + Staff & Service", "=B5+B6", False),
        ("Check", '=IF(B4="","Type your season total",IF(B7<=B4,"Within the total","Over by $"&TEXT(B7-B4,"0.00")&": lower a ceiling, or move a name to card-and-a-note"))', False),
        ("Spent so far (gifts + Staff & Service given)", "=SUM(I12:I51)+'Staff & Service'!C18", False)]
for i, (lab, f, is_in) in enumerate(labs, start=4):
    gl.cell(row=i, column=1, value=lab); cell = gl.cell(row=i, column=2, value=f)
    (inp if is_in else out)(cell, USD if i != 8 else None)
header(gl, 11, ["Name", "Ceiling", "Idea", "Test (U / E / K)", "Bought", "Wrapped", "Sent", "Thanked", "Spent", "Over / under"])
for r in range(12, 52):
    for col in range(1, 10): inp(gl.cell(row=r, column=col))
    gl.cell(row=r, column=2).number_format = USD; gl.cell(row=r, column=9).number_format = USD
    out(gl.cell(row=r, column=10, value=f'=IF(OR(I{r}="",B{r}=""),"",I{r}-B{r})'), USD)
listval(gl, "D12:D51", ["U", "E", "K", "U+E", "U+K", "E+K", "None: choose again"])
ticklist(gl, "E12:H51")
gl.conditional_formatting.add("A12:J51", FormulaRule(formula=['AND(ISNUMBER($I12),ISNUMBER($B12),$I12>$B12)'], fill=AMB))
widths(gl, {"A": 44, "B": 14, "C": 26, "D": 14, "E": 8, "F": 9, "G": 7, "H": 9, "I": 11, "J": 12})

# 9 Staff & Service
ss = sheet("Staff & Service", "Staff & Service First", "Typical guidance (Emily Post Institute), not law. Only the USPS line is a rule. Done by 15 December.")
ss["A3"] = "Days left until 15 December:"; ss["C3"] = '=MAX(0,DATE(YEAR(TODAY()),12,15)-TODAY())'; ss["C3"].font = Font(bold=True)
header(ss, 4, ["Who", "Typical guidance", "My sum $", "Card", "Given", "Date given"])
who = [("Cleaner or housekeeper", "Up to one week's pay"), ("Regular sitter or nanny", "Up to one week's pay, plus a small gift from the children"),
       ("Dog walker", "Up to one week's pay"), ("Hairdresser or barber", "The cost of one visit"), ("Teacher", "A gift, not cash (check the school's policy)"),
       ("US mail carrier", "USPS rule: a non-cash gift worth $20 or less; never cash or cash-like gift cards"), ("Doorman or building super", "$25–$100+, by city and effort"),
       ("Bin (trash) collectors", "$10–$30 each, where allowed"), ("Regular delivery driver", "Snacks or a small gift"), ("Carer for a relative", "A card and a gift; check the agency's policy"),
       ("", ""), ("", "")]
for i, (w, gd) in enumerate(who, start=5):
    inp(ss.cell(row=i, column=1, value=w or None)); ss.cell(row=i, column=2, value=gd or None).border = BOX
    inp(ss.cell(row=i, column=3), USD); inp(ss.cell(row=i, column=4)); inp(ss.cell(row=i, column=5)); inp(ss.cell(row=i, column=6), DATE)
ticklist(ss, "D5:E16")
ss["A17"] = "Total planned"; out(ss["C17"], USD, True); ss["C17"] = "=SUM(C5:C16)"
ss["A18"] = "Total given"; out(ss["C18"], USD, True); ss["C18"] = '=SUMIF(E5:E16,"' + TICK + '",C5:C16)'
widths(ss, {"A": 28, "B": 60, "C": 11, "D": 7, "E": 7, "F": 16})

# 10 Thaw Sum
t = sheet("Thaw Sum", "The Thaw Sum (fridge thaw)", "Use one guidance only. US (USDA): count back from the day, plus one spare day. UK (FSA): count back from the oven time, no spare day.")
tl = [("Guidance I follow (US or UK)", None, True, None), ("Bird's weight, lb (US)", None, True, "0.0"), ("Bird's weight, kg (UK)", None, True, "0.0"),
      ("T-0, the day (from At a Glance)", "='At a Glance'!B13", False, DATE), ("UK only: the date and time it goes in the oven", None, True, DT)]
for i, (lab, f, is_in, nf) in enumerate(tl, start=4):
    t.cell(row=i, column=1, value=lab); (inp if is_in else out)(t.cell(row=i, column=2, value=f), nf)
listval(t, "B4", ["US", "UK"])
t["A10"] = "US (USDA): about 24 hours per 4–5 lb; a fridge-thawed bird keeps 1–2 days before cooking"; t["A10"].font = Font(bold=True, color=INK)
t["A11"] = "Hours needed (lb ÷ 4 × 24)"; out(t["B11"], "0"); t["B11"] = '=IF(B5="","",B5/4*24)'
t["A12"] = "Days back: hours ÷ 24 rounded up, plus one spare day"; out(t["B12"], "0"); t["B12"] = '=IF(B11="","",ROUNDUP(B11/24,0)+1)'
t["A13"] = "INTO THE FRIDGE ON"; t["A13"].font = Font(bold=True); out(t["B13"], DATE, True); t["B13"] = '=IF(OR(B12="",B7=""),"",B7-B12)'
t["A15"] = "UK (FSA): about 10–12 hours per kg at 4°C; cook within 24 hours of fully defrosting"; t["A15"].font = Font(bold=True, color=INK)
t["A16"] = "Hours needed (kg × 12)"; out(t["B16"], "0"); t["B16"] = '=IF(B6="","",B6*12)'
t["A17"] = "INTO THE FRIDGE AT (oven time minus those hours; no spare day)"; t["A17"].font = Font(bold=True); out(t["B17"], DT, True); t["B17"] = '=IF(OR(B16="",B8=""),"",B8-B16/24)'
t["A18"] = "Earliest it could be fully defrosted (at 10 h per kg)"; out(t["B18"], DT); t["B18"] = '=IF(B17="","",B17+B6*10/24)'
t["A19"] = "Must be cooked by (24 hours after that)"; out(t["B19"], DT); t["B19"] = '=IF(B18="","",B18+1)'
t["A20"] = '=IF(AND(B6<>"",B6>12),"Over 12 kg: at the faster rate it could be defrosted more than 24 hours before the oven time. Check the FSA\'s current guidance.","")'
t["A22"] = "Sources: USDA FSIS (www.fsis.usda.gov); Food Standards Agency (food.gov.uk). If this ever differs from the regulator's current page, the regulator is right."
t["A22"].font = Font(italic=True, color="3A3B3D")
t.conditional_formatting.add("A20", FormulaRule(formula=['$A$20<>""'], fill=AMB))
widths(t, {"A": 70, "B": 26})

# 11 Oven (fill)
o["A1"] = "The Oven Timetable"; o["A1"].font = Font(bold=True, size=14, color=INK)
o["A2"] = "Two oven dishes whose times overlap are flagged; so is any guest's dish (Assignments tab) that needs the oven at the same time."; o["A2"].font = Font(italic=True, color="3A3B3D")
header(o, 4, ["Dish", "Appliance", "Temp (°F / °C)", "From", "To", "Clash?"])
for r in range(5, 35):
    for col in range(1, 6): inp(o.cell(row=r, column=col))
    o.cell(row=r, column=4).number_format = TIME; o.cell(row=r, column=5).number_format = TIME
    out(o.cell(row=r, column=6, value=(f'=IF(AND(B{r}="Oven",D{r}<>"",E{r}<>""),IF(SUMPRODUCT(($B$5:$B$34="Oven")*($D$5:$D$34<E{r})*($E$5:$E$34>D{r}))'
        f'+SUMPRODUCT((Assignments!$E$5:$E$24="Oven")*(Assignments!$F$5:$F$24<E{r})*(Assignments!$G$5:$G$24>D{r}))>1,'
        f'"Shares the oven: same temperature, and room for both?",""),"")')))
listval(o, "B5:B34", ["Oven", "Hob (stovetop)", "Microwave", "Slow cooker", "Cold / room temp"])
o.conditional_formatting.add("F5:F34", FormulaRule(formula=['$F5<>""'], fill=AMB))
widths(o, {"A": 28, "B": 18, "C": 14, "D": 11, "E": 11, "F": 44})

# 12 Thank-Yous
ty = sheet("Thank-Yous", "Thank-You by the Second", "Four lines: the thing, one specific detail, the person, the sign-off. Templates 1–6 are on page 37.")
ty["A3"] = "Deadline"; out(ty["B3"], DATE, True)
ty["B3"] = "=IF('At a Glance'!B13=\"\",\"\",IF(MONTH('At a Glance'!B13)=12,DATE(YEAR('At a Glance'!B13)+1,1,2),'At a Glance'!B13+2))"
ty["C3"] = "Days left"; ty["D3"] = '=IF(B3="","",MAX(0,B3-TODAY()))'
ty["E3"] = "Outstanding"; ty["F3"] = '=COUNTA(A6:A45)-COUNTIF(E6:E45,"' + TICK + '")'
header(ty, 5, ["Name", "For", "Template (1–6)", "Written", "Posted / given"])
for r in range(6, 46):
    for col in range(1, 6): inp(ty.cell(row=r, column=col))
listval(ty, "C6:C45", ["1", "2", "3", "4", "5", "6"]); ticklist(ty, "D6:E45")
widths(ty, {"A": 24, "B": 30, "C": 14, "D": 10, "E": 14, "F": 8})

# 13 Reckoning
rk = sheet("Reckoning", "The Reckoning (T+7)", "Planned against actual. Then type next year's T-0 date (B27): the Clock rewrites itself from it.")
header(rk, 4, ["Measure", "Planned", "Actual", "Difference"])
rl = [("Guests", "='At a Glance'!B9", None, "0"),
      ("Total spend: food, drink, table", "='At a Glance'!B10", "='Spend Log'!J12", USD),
      ("Per plate", "='At a Glance'!B11", '=IFERROR(C6/C5,"")', USD),
      ("Gift season total (incl. Staff & Service)", "='Gift Ledger'!B4", "='Gift Ledger'!B9", USD),
      ("Clock ticks (out of 25)", 25, '=COUNTIF(Clock!G6:G30,"' + TICK + '")', "0"),
      ("Minutes on the clock (ticks × 15)", 375, "=C9*15", "0"),
      ("Tasks done on their own day", None, None, "0"),
      ("Borrowed items returned within two days", None, None, "0"),
      ("Thank-yous sent by the deadline", None, None, "0")]
for i, (lab, pl, ac, nf) in enumerate(rl, start=5):
    rk.cell(row=i, column=1, value=lab).border = BOX
    cpl = rk.cell(row=i, column=2, value=pl); cac = rk.cell(row=i, column=3, value=ac)
    (out if isinstance(pl, str) or isinstance(pl, int) else inp)(cpl, nf)
    (out if isinstance(ac, str) else inp)(cac, nf)
    out(rk.cell(row=i, column=4, value=f'=IF(OR(B{i}="",C{i}=""),"",C{i}-B{i})'), nf)
rk["A15"] = "The calm line (1 = dread, 5 = composed)"; rk["A15"].font = Font(bold=True, color=INK)
header(rk, 16, ["When", "Calm score"])
for i, w in enumerate(["T-21", "T-15", "T-8", "T-1", "At the door", "T+1"], start=17):
    rk.cell(row=i, column=1, value=w).border = BOX; inp(rk.cell(row=i, column=2))
dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True); rk.add_data_validation(dv); dv.add("B17:B22")
ch = LineChart(); ch.title = "The calm line"; ch.y_axis.title = "Calm"; ch.y_axis.scaling.min = 1; ch.y_axis.scaling.max = 5; ch.height = 6; ch.width = 14
ch.add_data(Reference(rk, min_col=2, min_row=16, max_row=22), titles_from_data=True); ch.set_categories(Reference(rk, min_col=1, min_row=17, max_row=22))
rk.add_chart(ch, "F15")
rk["A25"] = "Next year"; rk["A25"].font = Font(bold=True, color=INK)
rk["A26"] = "Next year's starting Per-Plate Number"; inp(rk["B26"], USD)
rk["A27"] = "Next year's T-0 date (the Clock switches to it)"; inp(rk["B27"], DATE)
rk["A28"] = "Thanksgiving 2027: Thu 25 Nov 2027 · Christmas 2027: Sat 25 Dec 2027"; rk["A28"].font = Font(italic=True, color="3A3B3D")
rk["A30"] = "Three to keep"; rk["A31"] = "Three to change"; rk["A32"] = "One to drop entirely"
for r in (30, 31, 32): inp(rk.cell(row=r, column=2)); rk.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
widths(rk, {"A": 46, "B": 16, "C": 16, "D": 14})

order = ["At a Glance", "Clock", "Per-Plate", "Spend Log", "Assignments", "Drinks", "Shopping", "Gift Ledger",
         "Staff & Service", "Thaw Sum", "Oven", "Thank-Yous", "Reckoning"]
wb._sheets = [wb[n] for n in order]
assert len(wb.sheetnames) == 13
wb.save(OUT)
print("wrote", OUT, wb.sheetnames)
