"""Builds the free Gift Ledger sheet (gift-ledger-sheet.xlsx) for the lead magnet: tabs Drawer, Ledger, Tipping.
Run: python3 make_gift_ledger.py [out.xlsx]. Uploaded to Drive with conversion to make the Google template (build-spec.md, Job J2)."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.formatting.rule import FormulaRule
import make_planner_styles as st

OUT = sys.argv[1] if len(sys.argv) > 1 else "gift-ledger-sheet.xlsx"
wb = Workbook(); T = st.TICK; USD = st.USD

d = st.sheet(wb, "Drawer", "The Gift Drawer Count", "Count before you buy. Good / Mendable / Gone, then only the gap.", first=True)
st.header(d, 4, ["Category", "Count", "Good", "Mendable", "Gone", "Earmarked for", "Par (how many you need)", "Gap to buy"])
cats = ["Unopened gifts you won't use", "Gifts bought early and forgotten", "Candles (unused)", "Greeting cards and envelopes", "Stamps",
        "Wrapping paper (rolls)", "Ribbon, string, tags", "Tape (rolls)", "Gift bags and boxes", "Jars, bottles, tins", "Gift cards and vouchers", "Other"]
for i, c in enumerate(cats, start=5):
    st.inp(d.cell(row=i, column=1, value=c))
    for col in range(2, 8): st.inp(d.cell(row=i, column=col))
    st.out(d.cell(row=i, column=8, value=f'=IF(G{i}="","",MAX(0,G{i}-N(C{i})-N(D{i})))'))
d["A18"] = "Items to buy, in total"; d["H18"] = "=SUM(H5:H16)"; d["H18"].font = Font(bold=True)
st.widths(d, {"A": 34, "B": 8, "C": 8, "D": 10, "E": 8, "F": 22, "G": 12, "H": 11})

l = st.sheet(wb, "Ledger", "The Ledger", "Every name has a ceiling. Every idea passes Use / Eat / Keep ten years. The last column is Thanked.")
rows = [("Season gift total (your limit)", None, True), ("Staff & Service (from Tipping)", "=Tipping!C16", False),
        ("Ceilings, all names", "=SUM(B12:B41)", False), ("Ceilings + Staff & Service", "=B5+B6", False),
        ("Check", '=IF(B4="","Type your season total",IF(B7<=B4,"Within the total","Over by $"&TEXT(B7-B4,"0.00")))', False),
        ("Spent so far", "=SUM(I12:I41)+Tipping!C17", False)]
for i, (lab, f, is_in) in enumerate(rows, start=4):
    l.cell(row=i, column=1, value=lab); (st.inp if is_in else st.out)(l.cell(row=i, column=2, value=f), USD if i != 8 else None)
st.header(l, 11, ["Name", "Gift Ceiling", "Idea", "Test (U / E / K)", "Bought", "Wrapped", "Sent", "Thanked", "Spent", "Over / under"])
for r in range(12, 42):
    for col in range(1, 10): st.inp(l.cell(row=r, column=col))
    l.cell(row=r, column=2).number_format = USD; l.cell(row=r, column=9).number_format = USD
    st.out(l.cell(row=r, column=10, value=f'=IF(OR(I{r}="",B{r}=""),"",I{r}-B{r})'), USD)
st.listval(l, "D12:D41", ["U", "E", "K", "U+E", "U+K", "E+K", "None: choose again"]); st.ticklist(l, "E12:H41")
l.conditional_formatting.add("A12:J41", FormulaRule(formula=['AND(ISNUMBER($I12),ISNUMBER($B12),$I12>$B12)'], fill=st.AMB))
st.widths(l, {"A": 32, "B": 13, "C": 26, "D": 14, "E": 8, "F": 9, "G": 7, "H": 9, "I": 11, "J": 12})

t = st.sheet(wb, "Tipping", "The Holiday Tipping Card", "Typical guidance (Emily Post Institute), not law. Only the USPS line is a rule. Done by 15 December.")
t["A3"] = "Days left until 15 December:"; t["C3"] = '=MAX(0,DATE(YEAR(TODAY()),12,15)-TODAY())'; t["C3"].font = Font(bold=True)
st.header(t, 4, ["Who", "Typical guidance", "My sum $", "Card written", "Given", "Date"])
who = [("Cleaner or housekeeper", "Up to one week's pay"), ("Regular sitter or nanny", "Up to one week's pay, plus a small gift from the children"),
       ("Dog walker", "Up to one week's pay"), ("Hairdresser or barber", "The cost of one visit"), ("Teacher", "A gift, not cash (check the school's policy first)"),
       ("US mail carrier", "USPS rule: a non-cash gift worth $20 or less; never cash or cash-like gift cards"), ("Doorman or building super", "$25–$100+, by city and effort"),
       ("Bin (trash) collectors", "$10–$30 each, where allowed"), ("Regular delivery driver", "Snacks or a small gift"), ("", ""), ("", "")]
for i, (w, g) in enumerate(who, start=5):
    st.inp(t.cell(row=i, column=1, value=w or None)); t.cell(row=i, column=2, value=g or None).border = st.BOX
    st.inp(t.cell(row=i, column=3), USD); st.inp(t.cell(row=i, column=4)); st.inp(t.cell(row=i, column=5)); st.inp(t.cell(row=i, column=6), st.DATE)
st.ticklist(t, "D5:E15")
t["A16"] = "Total for Staff & Service (flows into the Ledger)"; st.out(t["C16"], USD, True); t["C16"] = "=SUM(C5:C15)"
t["A17"] = "Total given"; st.out(t["C17"], USD, True); t["C17"] = f'=SUMIF(E5:E15,"{T}",C5:C15)'
st.widths(t, {"A": 30, "B": 62, "C": 11, "D": 12, "E": 7, "F": 14})
wb.save(OUT); print("wrote", OUT, wb.sheetnames)
