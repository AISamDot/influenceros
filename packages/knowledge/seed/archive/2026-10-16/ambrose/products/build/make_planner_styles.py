"""Shared styles for the sheet builders."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
INK, MID, SILVER, LINEN, AMBER = "1F3440", "161D2A", "C9CCCF", "F1EEE7", "D9A55B"
HDR = PatternFill("solid", fgColor=INK); INP = PatternFill("solid", fgColor=LINEN); AMB = PatternFill("solid", fgColor=AMBER)
thin = Side(style="thin", color=SILVER); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
TICK = "✓"; USD = '"$"#,##0.00'; DATE = "ddd d mmm yyyy"
def sheet(wb, name, title, note=None, first=False):
    ws = wb.active if first else wb.create_sheet(); ws.title = name
    ws["A1"] = title; ws["A1"].font = Font(bold=True, size=14, color=INK)
    if note: ws["A2"] = note; ws["A2"].font = Font(italic=True, color="3A3B3D")
    ws.sheet_view.showGridLines = False; return ws
def header(ws, row, labels, col=1):
    for i, l in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=l); c.fill = HDR; c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BOX
def inp(c, fmt=None):
    c.fill = INP; c.border = BOX
    if fmt: c.number_format = fmt
def out(c, fmt=None, bold=False):
    c.border = BOX; c.font = Font(bold=bold, color=MID)
    if fmt: c.number_format = fmt
def widths(ws, w):
    for k, v in w.items(): ws.column_dimensions[k].width = v
def ticklist(ws, rng):
    dv = DataValidation(type="list", formula1=f'"{TICK}"', allow_blank=True); ws.add_data_validation(dv); dv.add(rng)
def listval(ws, rng, items):
    dv = DataValidation(type="list", formula1='"' + ",".join(items) + '"', allow_blank=True); ws.add_data_validation(dv); dv.add(rng)
